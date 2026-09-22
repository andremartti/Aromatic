# -*- coding: utf-8 -*-
"""
AUDITORÍA AUTOMÁTICA del sistema financiero AROMATIC.

Qué hace:
  1. Copia el libro generado.
  2. Inyecta un juego de datos de prueba que recorre TODO el flujo:
     producto → costo → compra → inventario → venta → utilidad → gasto → caja
     → rentabilidad → metas → equilibrio → escenarios → dashboard.
     Incluye casos extremos: sin precio, sin costo, precio = costo, precio < costo,
     venta sin inventario, venta anulada, SKU duplicado, gasto sin categoría,
     compra sin proveedor, cobros y pagos parciales.
  3. Recalcula el libro con LibreOffice (motor de cálculo real).
  4. Compara cada resultado contra el valor esperado calculado a mano.

Uso:  python3 generador/validar_sistema.py
"""
import datetime as dt
import os
import shutil
import subprocess
import sys

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEN = os.path.join(BASE, "AROMATIC_Sistema_Financiero.xlsx")
TMP = os.environ.get("AROMATIC_TMP", "/tmp/aromatic_val")
PRUEBA = os.path.join(TMP, "prueba.xlsx")
SALIDA = os.path.join(TMP, "out")

# --- Coordenadas (deben coincidir con construir_sistema.py) -----------------
R_PROD, R_COST, R_COMPRAS, R_VENTAS, R_GASTOS, R_INV = 7, 8, 7, 7, 7, 7
R_CLI, R_PROV, R_CAJAMOV = 7, 7, 7
MOV_COL = 21                      # columna U de CAJA
RENT_R, RENT_TOT = 8, 44
RENT_R_PROD = 8
CAJA_R, CAJA_TOT = 7, 43
MET_R = 7
CTRL_R = 8

F = dt.datetime
d = lambda y, m, dd: dt.datetime(y, m, dd)  # noqa: E731


def cargar():
    os.makedirs(TMP, exist_ok=True)
    shutil.rmtree(SALIDA, ignore_errors=True)
    os.makedirs(SALIDA, exist_ok=True)
    shutil.copy(ORIGEN, PRUEBA)
    return openpyxl.load_workbook(PRUEBA)


def poblar(wb):
    cfg, pro, cos = wb["CONFIG"], wb["PRODUCTOS"], wb["COSTOS"]
    com, ven, gas = wb["COMPRAS"], wb["VENTAS"], wb["GASTOS"]
    caj, cli, prv = wb["CAJA"], wb["CLIENTES"], wb["PROVEEDORES"]
    eq, met = wb["EQUILIBRIO"], wb["METAS"]

    # ---- CONFIG
    cfg["B8"] = d(2026, 9, 1)      # inicio del calendario
    cfg["B10"] = 5000              # saldo inicial de caja
    cfg["B11"] = 20000             # capital aportado

    # ---- COSTOS (fila 8 = GEL-250, 9 = GEL-500, 10 = GEL-2L, 11 = GEL-GAL,
    #              12 = JAB-500, 13 = JAB-2L)
    cos["D8"], cos["F8"], cos["G8"], cos["I8"] = 20, 1, 2, 1    # costo real = 24
    cos["P8"], cos["R8"] = 0.05, 3                              # 5% comisión + L3 delivery
    cos["D9"] = 80                                              # GEL-500 costo 80 (> precio)
    cos["D11"] = 100                                            # GEL-GAL con costo, sin precio
    cos["D12"] = 60                                             # JAB-500 costo = precio
    cos["D13"] = 60                                             # JAB-2L costo estándar

    # ---- PRODUCTOS
    pro["K7"], pro["Y7"], pro["Z7"] = 35, 10, 50                # precio normal
    pro["K8"] = 70                                              # precio < costo  → margen negativo
    pro["K9"] = 140                                             # precio sin costo → FALTA COSTO
    pro["K11"] = 60                                             # precio = costo
    pro["K12"] = 100                                            # se venderá sin inventario
    pro["A66"], pro["B66"] = "GEL-250", "SKU duplicado de prueba"

    # ---- CLIENTES / PROVEEDORES
    cli["B7"], cli["C7"], cli["E7"] = "Cliente Test", "9999-9999", "WhatsApp"
    prv["B7"], prv["G7"] = "Proveedor Test", "Contado"

    # ---- COMPRAS
    r = R_COMPRAS
    com[f"B{r}"], com[f"C{r}"], com[f"D{r}"] = d(2026, 9, 5), "Proveedor Test", "GEL-250"
    com[f"G{r}"], com[f"H{r}"], com[f"J{r}"], com[f"K{r}"] = 100, 20, 150, 50
    com[f"O{r}"], com[f"P{r}"], com[f"Q{r}"], com[f"R{r}"] = "Efectivo", "Pagado", d(2026, 9, 5), 2200
    r += 1
    com[f"B{r}"], com[f"D{r}"] = d(2026, 9, 20), "GEL-500"      # sin proveedor (caso extremo)
    com[f"G{r}"], com[f"H{r}"], com[f"P{r}"] = 50, 80, "Pendiente"

    # ---- VENTAS
    r = R_VENTAS
    ven[f"B{r}"], ven[f"C{r}"], ven[f"D{r}"], ven[f"E{r}"] = d(2026, 9, 10), "F-001", "Cliente Test", "WhatsApp"
    ven[f"F{r}"], ven[f"I{r}"], ven[f"L{r}"] = "GEL-250", 10, 20
    ven[f"S{r}"], ven[f"T{r}"], ven[f"U{r}"], ven[f"V{r}"] = "Efectivo", "Cobrado", d(2026, 9, 10), 330
    r += 1
    ven[f"B{r}"], ven[f"C{r}"], ven[f"D{r}"], ven[f"E{r}"] = d(2026, 9, 18), "F-002", "Cliente Test", "Facebook"
    ven[f"F{r}"], ven[f"I{r}"], ven[f"T{r}"] = "JAB-2L", 5, "Pendiente"   # venta sin inventario
    r += 1
    ven[f"B{r}"], ven[f"C{r}"], ven[f"D{r}"], ven[f"E{r}"] = d(2026, 9, 25), "F-003", "Cliente Test", "WhatsApp"
    ven[f"F{r}"], ven[f"I{r}"], ven[f"T{r}"] = "GEL-500", 3, "Anulado"    # venta anulada
    r += 1
    ven[f"B{r}"], ven[f"C{r}"], ven[f"D{r}"], ven[f"E{r}"] = d(2026, 10, 5), "F-004", "Cliente Test", "WhatsApp"
    ven[f"F{r}"], ven[f"I{r}"] = "GEL-250", 5
    ven[f"S{r}"], ven[f"T{r}"], ven[f"U{r}"], ven[f"V{r}"] = "Efectivo", "Cobrado", d(2026, 10, 5), 175

    # ---- GASTOS
    r = R_GASTOS
    gas[f"B{r}"], gas[f"C{r}"], gas[f"D{r}"], gas[f"E{r}"] = d(2026, 9, 15), "Fijo", "Alquiler", "Alquiler local"
    gas[f"G{r}"], gas[f"H{r}"], gas[f"I{r}"], gas[f"J{r}"], gas[f"K{r}"] = 1000, "Efectivo", "Pagado", d(2026, 9, 15), 1000
    r += 1
    gas[f"B{r}"], gas[f"C{r}"], gas[f"E{r}"] = d(2026, 9, 16), "Variable", "Gasto sin categoría"  # sin categoría
    gas[f"G{r}"], gas[f"I{r}"] = 500, "Pendiente"
    r += 1
    gas[f"B{r}"], gas[f"C{r}"], gas[f"D{r}"], gas[f"E{r}"] = d(2026, 9, 20), "Variable", "Publicidad", "Facebook Ads"
    gas[f"G{r}"], gas[f"I{r}"], gas[f"J{r}"], gas[f"K{r}"], gas[f"M{r}"] = 400, "Pagado", d(2026, 9, 20), 400, "Facebook"

    # ---- CAJA · movimiento manual (aporte de capital)
    r = R_CAJAMOV
    caj.cell(row=r, column=MOV_COL + 1, value=d(2026, 9, 1))
    caj.cell(row=r, column=MOV_COL + 2, value="Entrada")
    caj.cell(row=r, column=MOV_COL + 3, value="Aporte de capital")
    caj.cell(row=r, column=MOV_COL + 4, value="Aporte inicial del propietario")
    caj.cell(row=r, column=MOV_COL + 5, value=20000)
    caj.cell(row=r, column=MOV_COL + 6, value="Efectivo")

    # ---- EQUILIBRIO en modo manual (para probar ambas rutas)
    eq["B7"], eq["D8"], eq["D9"], eq["D10"] = "Manual", 1000, 50, 30
    eq["B26"] = 5000

    # ---- METAS del primer mes
    met[f"B{MET_R}"], met[f"C{MET_R}"], met[f"D{MET_R}"] = 1000, 200, 20

    wb.save(PRUEBA)


def recalcular():
    env = dict(os.environ, HOME="/root")
    cmd = ["soffice", "--headless", "--norestore", "--convert-to", "xlsx",
           "--outdir", SALIDA, PRUEBA]
    subprocess.run(cmd, env=env, capture_output=True, timeout=600)
    salida = os.path.join(SALIDA, "prueba.xlsx")
    if not os.path.exists(salida):
        print("ERROR: LibreOffice no generó el archivo recalculado")
        sys.exit(1)
    return openpyxl.load_workbook(salida, data_only=True)


# ----------------------------------------------------------------- Chequeos --
RES = {"ok": 0, "fail": 0}
FALLOS = []


def chk(nombre, obtenido, esperado, tol=0.01):
    ok = False
    if isinstance(esperado, dt.datetime):
        ok = (obtenido == esperado)
    elif isinstance(esperado, bool):
        ok = (obtenido is esperado)
    elif isinstance(esperado, str):
        ok = (str(obtenido).strip() == esperado)
    elif esperado is None:
        ok = obtenido is None or obtenido == "" 
    else:
        try:
            ok = abs(float(obtenido) - float(esperado)) <= tol
        except (TypeError, ValueError):
            ok = False
    RES["ok" if ok else "fail"] += 1
    marca = "  ok  " if ok else " FALLA"
    linea = f"[{marca}] {nombre:<62} obtenido={obtenido!r:<30} esperado={esperado!r}"
    if not ok:
        FALLOS.append(linea)
    print(linea)
    return ok


def main():
    print("1) Inyectando datos de prueba…")
    wb = cargar()
    poblar(wb)
    print("2) Recalculando con LibreOffice…")
    r = recalcular()
    p, c, i, v, cm, g = r["PRODUCTOS"], r["COSTOS"], r["INVENTARIO"], r["VENTAS"], r["COMPRAS"], r["GASTOS"]
    rn, ca, me, eq, es, ct, da = r["RENTABILIDAD"], r["CAJA"], r["METAS"], r["EQUILIBRIO"], r["ESCENARIOS"], r["CONTROL"], r["DASHBOARD"]
    cl, pv = r["CLIENTES"], r["PROVEEDORES"]
    print("3) Auditando resultados\n")

    print("── COSTOS ────────────────────────────────────────────")
    chk("COSTOS · costo real unitario GEL-250 (20+1+2+1)", c["O8"].value, 24)
    chk("COSTOS · costo variable de venta (35×5% + 3)", c["U8"].value, 4.75)
    chk("COSTOS · costo variable total (24 + 4.75)", c["V8"].value, 28.75)
    chk("COSTOS · sin costo capturado → PENDIENTE", c["O10"].value, "PENDIENTE")

    print("\n── PRODUCTOS ─────────────────────────────────────────")
    chk("PRODUCTOS · costo real unitario (desde COSTOS)", p["Q7"].value, 24)
    chk("PRODUCTOS · utilidad unitaria (35 − 24)", p["T7"].value, 11)
    chk("PRODUCTOS · margen bruto (11/35)", p["U7"].value, 0.314285, 0.0001)
    chk("PRODUCTOS · markup (11/24)", p["V7"].value, 0.458333, 0.0001)
    chk("PRODUCTOS · margen de contribución (35 − 28.75)", p["W7"].value, 6.25)
    chk("PRODUCTOS · estado de datos completo", p["AB7"].value, "COMPLETO")
    chk("PRODUCTOS · precio < costo → margen negativo", p["T8"].value, -10)
    chk("PRODUCTOS · estado MARGEN NEGATIVO", p["AB8"].value, "MARGEN NEGATIVO")
    chk("PRODUCTOS · precio sin costo → utilidad PENDIENTE", p["T9"].value, "PENDIENTE")
    chk("PRODUCTOS · estado FALTA COSTO", p["AB9"].value, "FALTA COSTO")
    chk("PRODUCTOS · costo sin precio → estado FALTA PRECIO", p["AB10"].value, "FALTA PRECIO")
    chk("PRODUCTOS · precio = costo → utilidad 0", p["T11"].value, 0)
    chk("PRODUCTOS · estado PRECIO = COSTO", p["AB11"].value, "PRECIO = COSTO")
    chk("PRODUCTOS · producto sin datos → utilidad PENDIENTE", p["T20"].value, "PENDIENTE")
    chk("PRODUCTOS · fila vacía no genera error", p["T60"].value, None)
    chk("PRODUCTOS · precio de referencia 2016 intacto (informativo)", p["J7"].value, 35)

    print("\n── COMPRAS ───────────────────────────────────────────")
    chk("COMPRAS · ID automático", cm["A7"].value, "C-0001")
    chk("COMPRAS · costo total (100×20)", cm["I7"].value, 2000)
    chk("COMPRAS · costo total de adquisición (+150+50)", cm["M7"].value, 2200)
    chk("COMPRAS · costo unitario de adquisición (2200/100)", cm["N7"].value, 22)
    chk("COMPRAS · saldo por pagar (pagada)", cm["S7"].value, 0)
    chk("COMPRAS · saldo por pagar (pendiente)", cm["S8"].value, 4000)
    chk("COMPRAS · alerta por falta de proveedor", "Falta proveedor" in str(cm["U8"].value), True)
    chk("COMPRAS · fila correcta sin alertas", cm["U7"].value, "OK")

    print("\n── INVENTARIO ────────────────────────────────────────")
    chk("INVENTARIO · compras GEL-250", i["F7"].value, 100)
    chk("INVENTARIO · ventas GEL-250 (10 + 5, anuladas fuera)", i["G7"].value, 15)
    chk("INVENTARIO · stock actual (0+100−15)", i["I7"].value, 85)
    chk("INVENTARIO · costo promedio ponderado (2200/100)", i["M7"].value, 22)
    chk("INVENTARIO · costo aplicado = WAC", i["N7"].value, 22)
    chk("INVENTARIO · método de costo", i["O7"].value, "Promedio ponderado")
    chk("INVENTARIO · valor del inventario (85×22)", i["P7"].value, 1870)
    chk("INVENTARIO · estado OK", i["S7"].value, "OK")
    chk("INVENTARIO · venta anulada no descuenta stock (GEL-500)", i["I8"].value, 50)
    chk("INVENTARIO · stock negativo detectado (JAB-2L)", i["I12"].value, -5)
    chk("INVENTARIO · estado ERROR por stock negativo", i["S12"].value, "ERROR")
    chk("INVENTARIO · costo estándar cuando no hay compras", i["N12"].value, 60)
    chk("INVENTARIO · método de costo estándar", i["O12"].value, "Costo estándar")
    chk("INVENTARIO · producto sin costo → PENDIENTE", i["O20"].value, "PENDIENTE")
    chk("INVENTARIO · valor PENDIENTE si no hay costo", i["P20"].value, "PENDIENTE")
    chk("INVENTARIO · última compra registrada", i["U7"].value, dt.datetime(2026, 9, 5))

    print("\n── VENTAS ────────────────────────────────────────────")
    chk("VENTAS · ID automático", v["A7"].value, "V-0001")
    chk("VENTAS · precio traído de PRODUCTOS", v["K7"].value, 35)
    chk("VENTAS · venta bruta (10×35)", v["M7"].value, 350)
    chk("VENTAS · venta neta (350−20)", v["N7"].value, 330)
    chk("VENTAS · costo unitario aplicado (WAC)", v["O7"].value, 22)
    chk("VENTAS · costo total (10×22)", v["P7"].value, 220)
    chk("VENTAS · utilidad bruta (330−220)", v["Q7"].value, 110)
    chk("VENTAS · margen (110/330)", v["R7"].value, 0.333333, 0.0001)
    chk("VENTAS · saldo por cobrar (cobrada)", v["W7"].value, 0)
    chk("VENTAS · saldo por cobrar (pendiente)", v["W8"].value, 500)
    chk("VENTAS · venta anulada → unidades efectivas 0", v["J9"].value, 0)
    chk("VENTAS · venta anulada → venta neta 0", v["N9"].value, 0)
    chk("VENTAS · alerta de stock negativo", "Stock negativo" in str(v["Y8"].value), True)
    chk("VENTAS · fila correcta sin alertas", v["Y7"].value, "OK")

    print("\n── GASTOS ────────────────────────────────────────────")
    chk("GASTOS · ID automático", g["A7"].value, "G-0001")
    chk("GASTOS · saldo por pagar", g["L8"].value, 500)
    chk("GASTOS · alerta SIN CATEGORÍA", "SIN CATEGORÍA" in str(g["O8"].value), True)

    print("\n── RENTABILIDAD (mes 1: sep-2026) ────────────────────")
    chk("RENT · ventas brutas (350+500+0)", rn[f"B{RENT_R}"].value, 850)
    chk("RENT · descuentos", rn[f"C{RENT_R}"].value, 20)
    chk("RENT · ventas netas (330+500)", rn[f"D{RENT_R}"].value, 830)
    chk("RENT · unidades (10+5)", rn[f"E{RENT_R}"].value, 15)
    chk("RENT · costo de ventas (220+300)", rn[f"F{RENT_R}"].value, 520)
    chk("RENT · utilidad bruta (830−520)", rn[f"G{RENT_R}"].value, 310)
    chk("RENT · margen bruto", rn[f"H{RENT_R}"].value, 310 / 830, 0.0001)
    chk("RENT · gastos fijos", rn[f"I{RENT_R}"].value, 1000)
    chk("RENT · gastos variables (500+400)", rn[f"J{RENT_R}"].value, 900)
    chk("RENT · utilidad operativa (310−1900)", rn[f"L{RENT_R}"].value, -1590)
    chk("RENT · aporte de capital NO es ingreso", rn[f"M{RENT_R}"].value, 0)
    chk("RENT · utilidad neta", rn[f"O{RENT_R}"].value, -1590)
    chk("RENT · transacciones del mes", rn[f"Q{RENT_R}"].value, 3)
    chk("RENT · ticket promedio (830/3)", rn[f"R{RENT_R}"].value, 830 / 3, 0.01)
    chk("RENT · publicidad del mes", rn[f"T{RENT_R}"].value, 400)
    chk("RENT · ROAS (830/400)", rn[f"U{RENT_R}"].value, 2.075, 0.001)
    chk("RENT · compras del mes (2200+4000)", rn[f"V{RENT_R}"].value, 6200)
    chk("RENT · ROI del mes", rn[f"W{RENT_R}"].value, -1590 / 8100, 0.0001)
    print("── RENTABILIDAD (mes 2: oct-2026) ────────────────────")
    chk("RENT · ventas netas mes 2", rn[f"D{RENT_R + 1}"].value, 175)
    chk("RENT · costo de ventas mes 2 (5×22)", rn[f"F{RENT_R + 1}"].value, 110)
    chk("RENT · utilidad bruta mes 2", rn[f"G{RENT_R + 1}"].value, 65)
    chk("RENT · ventas acumuladas mes 2", rn[f"X{RENT_R + 1}"].value, 1005)
    chk("RENT · total ventas netas", rn[f"D{RENT_TOT}"].value, 1005)
    chk("RENT · total utilidad neta", rn[f"O{RENT_TOT}"].value, -1525)
    print("── RENTABILIDAD (por producto / categoría / canal) ────")
    chk("RENT · GEL-250 unidades", rn[f"AD{RENT_R_PROD}"].value, 15)
    chk("RENT · GEL-250 ventas netas", rn[f"AE{RENT_R_PROD}"].value, 505)
    chk("RENT · GEL-250 utilidad bruta", rn[f"AG{RENT_R_PROD}"].value, 175)
    chk("RENT · categoría Hogar ventas (todas las ventas)", rn["AT8"].value, 1005)
    chk("RENT · canal WhatsApp ventas (330+175)", rn["BB9"].value, 505)
    chk("RENT · canal Facebook ROAS (500/400)", rn["BH10"].value, 1.25, 0.001)
    print("── RENTABILIDAD (resumen acumulado) ──────────────────")
    chk("RENT · inversión total acumulada (6200+1900)", rn["BK25"].value, 8100)
    chk("RENT · ROI sobre la inversión", rn["BK26"].value, -1525 / 8100, 0.0001)
    chk("RENT · capital aportado (20000 config + 20000 caja)", rn["BK27"].value, 40000)
    chk("RENT · valor del inventario (incluye SKU duplicado)", rn["BK29"].value, 7440)
    chk("RENT · cuentas por cobrar", rn["BK30"].value, 500)
    chk("RENT · cuentas por pagar (4000 + 500)", rn["BK31"].value, 4500)
    chk("RENT · capital de trabajo comprometido", rn["BK32"].value, 7440 + 500 - 4500)

    print("\n── CAJA ──────────────────────────────────────────────")
    chk("CAJA · saldo inicial (config)", ca[f"B{CAJA_R}"].value, 5000)
    chk("CAJA · cobros de clientes mes 1", ca[f"C{CAJA_R}"].value, 330)
    chk("CAJA · aportes de capital mes 1", ca[f"D{CAJA_R}"].value, 20000)
    chk("CAJA · total entradas", ca[f"G{CAJA_R}"].value, 20330)
    chk("CAJA · pagos a proveedores", ca[f"H{CAJA_R}"].value, 2200)
    chk("CAJA · pagos de gastos (1000+400)", ca[f"I{CAJA_R}"].value, 1400)
    chk("CAJA · total salidas", ca[f"M{CAJA_R}"].value, 3600)
    chk("CAJA · flujo neto del mes", ca[f"N{CAJA_R}"].value, 16730)
    chk("CAJA · saldo final", ca[f"O{CAJA_R}"].value, 21730)
    chk("CAJA · utilidad neta del mes (≠ flujo)", ca[f"P{CAJA_R}"].value, -1590)
    chk("CAJA · diferencia utilidad − flujo", ca[f"Q{CAJA_R}"].value, -18320)
    chk("CAJA · cuentas por cobrar al cierre", ca[f"R{CAJA_R}"].value, 500)
    chk("CAJA · cuentas por pagar al cierre", ca[f"S{CAJA_R}"].value, 4500)
    chk("CAJA · saldo inicial mes 2 = saldo final mes 1", ca[f"B{CAJA_R + 1}"].value, 21730)
    chk("CAJA · saldo final mes 2", ca[f"O{CAJA_R + 1}"].value, 21905)
    chk("CAJA · saldo de caja actual (indicador)", ca[f"E{CAJA_TOT + 3}"].value, 21905)

    print("\n── CLIENTES / PROVEEDORES ────────────────────────────")
    chk("CLIENTES · número de compras", cl["I7"].value, 4)
    chk("CLIENTES · ventas acumuladas", cl["J7"].value, 1005)
    chk("CLIENTES · saldo por cobrar", cl["M7"].value, 500)
    chk("CLIENTES · primera compra", cl["G7"].value, dt.datetime(2026, 9, 10))
    chk("CLIENTES · última compra", cl["H7"].value, dt.datetime(2026, 10, 5))
    chk("PROVEEDORES · compras acumuladas", pv["J7"].value, 2200)
    chk("PROVEEDORES · compra promedio", pv["L7"].value, 2200)
    chk("PROVEEDORES · saldo por pagar", pv["N7"].value, 0)

    print("\n── EQUILIBRIO ────────────────────────────────────────")
    chk("EQUILIBRIO · costos fijos automáticos", eq["C8"].value, 1000)
    chk("EQUILIBRIO · precio promedio automático (1005/20)", eq["C9"].value, 50.25)
    chk("EQUILIBRIO · costo variable unitario automático", eq["C10"].value, (630 + 900) / 20)
    chk("EQUILIBRIO · valor usado en modo Manual", eq["B8"].value, 1000)
    chk("EQUILIBRIO · margen de contribución (50−30)", eq["B14"].value, 20)
    chk("EQUILIBRIO · margen de contribución %", eq["B15"].value, 0.4, 0.0001)
    chk("EQUILIBRIO · punto de equilibrio unidades (1000/20)", eq["B16"].value, 50)
    chk("EQUILIBRIO · punto de equilibrio en L (50×50)", eq["B17"].value, 2500)
    chk("EQUILIBRIO · unidades promedio reales (20/2)", eq["B20"].value, 10)
    chk("EQUILIBRIO · margen de seguridad negativo", eq["B21"].value, -40)
    chk("EQUILIBRIO · simulador unidades para ganar 5000", eq["B27"].value, 300)
    chk("EQUILIBRIO · simulador ventas necesarias", eq["B28"].value, 15000)
    chk("EQUILIBRIO · sensibilidad precio −10% / costo −10%", eq["H33"].value, 56)

    print("\n── ESCENARIOS ────────────────────────────────────────")
    chk("ESCENARIOS · unidades base (20/2 meses)", es["B7"].value, 10)
    chk("ESCENARIOS · precio base", es["B8"].value, 50.25)
    chk("ESCENARIOS · costo variable unitario base (630/20)", es["B9"].value, 31.5)
    chk("ESCENARIOS · gastos fijos base", es["B11"].value, 1000)
    chk("ESCENARIOS · ventas base", es["B18"].value, 502.5)
    chk("ESCENARIOS · margen de contribución base", es["B20"].value, 187.5)
    chk("ESCENARIOS · utilidad operativa base", es["B23"].value, -1262.5)
    chk("ESCENARIOS · punto de equilibrio base (unid.)", es["B25"].value, 78)
    chk("ESCENARIOS · escenario vacío → PENDIENTE", es["C18"].value, "PENDIENTE")

    print("\n── METAS ─────────────────────────────────────────────")
    chk("METAS · ventas reales del mes", me[f"G{MET_R}"].value, 830)
    chk("METAS · % cumplimiento de ventas", me[f"L{MET_R}"].value, 0.83, 0.0001)
    chk("METAS · diferencia contra la meta", me[f"P{MET_R}"].value, -170)
    chk("METAS · venta diaria requerida (1000/26)", me[f"S{MET_R}"].value, 1000 / 26, 0.01)
    chk("METAS · mes sin meta", me[f"W{MET_R + 5}"].value, "SIN META")

    print("\n── CONTROL ───────────────────────────────────────────")
    valores = {ct[f"B{CTRL_R + k}"].value: ct[f"D{CTRL_R + k}"].value for k in range(32)}
    chk("CONTROL · SKU duplicados detectados", valores.get("SKU duplicados en PRODUCTOS"), 2)
    chk("CONTROL · productos con margen negativo",
        valores.get("Productos con MARGEN NEGATIVO (vende por debajo del costo)"), 1)
    chk("CONTROL · productos sin costo (41 SKU − 6 con costo)", valores.get("Productos sin costo capturado"), 35)
    chk("CONTROL · stock negativo",
        valores.get("Productos con STOCK NEGATIVO (vendió más de lo que tenía)"), 1)
    chk("CONTROL · gastos sin categoría", valores.get("Gastos SIN CATEGORÍA"), 1)
    chk("CONTROL · compras sin proveedor", valores.get("Compras sin proveedor"), 1)
    chk("CONTROL · precio = costo", valores.get("Productos con precio igual al costo (utilidad cero)"), 1)
    chk("CONTROL · ventas pendientes de cobro", valores.get("Ventas pendientes de cobro (informativo)"), 1)
    chk("CONTROL · diagnóstico por SKU (sin « · » colgando)", ct["Q8"].value, "SKU DUPLICADO")
    chk("CONTROL · diagnóstico JAB-2L", "STOCK NEGATIVO" in str(ct["Q13"].value), True)

    print("\n── DASHBOARD ─────────────────────────────────────────")
    chk("DASHBOARD · ventas netas del mes", da["B8"].value, 830)
    chk("DASHBOARD · unidades del mes", da["E8"].value, 15)
    chk("DASHBOARD · utilidad bruta del mes", da["B12"].value, 310)
    chk("DASHBOARD · utilidad neta del mes", da["H12"].value, -1590)
    chk("DASHBOARD · ventas acumuladas", da["B16"].value, 1005)
    chk("DASHBOARD · valor del inventario", da["B20"].value, 7440)
    chk("DASHBOARD · entradas de efectivo del mes", da["B24"].value, 20330)
    chk("DASHBOARD · saldo al cierre del mes", da["K24"].value, 21730)
    chk("DASHBOARD · cumplimiento de ventas", da["B28"].value, 0.83, 0.0001)
    chk("DASHBOARD · punto de equilibrio (unid.)", da["B32"].value, 50)
    chk("DASHBOARD · top 1 por utilidad (JAB-2L, L200)", da["P36"].value, "Jabón Líquido para Manos · 2 Litros")
    chk("DASHBOARD · utilidad del top 1", da["Q36"].value, 200)

    print("\n── FILAS VACÍAS (no deben generar ID ni alertas) ─────")
    chk("COMPRAS · fila vacía sin ID", cm["A50"].value, None)
    chk("COMPRAS · fila vacía sin alerta", cm["U50"].value, None)
    chk("VENTAS · fila vacía sin ID", v["A50"].value, None)
    chk("VENTAS · fila vacía sin alerta", v["Y50"].value, None)
    chk("GASTOS · fila vacía sin ID", g["A50"].value, None)
    chk("GASTOS · fila vacía sin alerta", g["O50"].value, None)
    chk("CAJA · movimiento vacío sin ID", r["CAJA"].cell(row=50, column=MOV_COL).value, None)
    chk("CLIENTES · fila vacía sin ID", cl["A50"].value, None)
    chk("INVENTARIO · fila sin producto vacía", i["A60"].value, None)

    print("\n── ERRORES DE EXCEL EN TODO EL LIBRO ─────────────────")
    errores = []
    for hoja in r.sheetnames:
        h = r[hoja]
        for fila in h.iter_rows():
            for cel in fila:
                if isinstance(cel.value, str) and cel.value.startswith("#") and cel.value.endswith(("!", "?", "0!", "A", "E")):
                    if cel.value in ("#DIV/0!", "#N/A", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!"):
                        errores.append(f"{hoja}!{cel.coordinate} = {cel.value}")
    chk("Ninguna celda con error de Excel", len(errores), 0)
    for e in errores[:40]:
        print("    ", e)

    print("\n" + "=" * 90)
    print(f"RESULTADO: {RES['ok']} pruebas correctas · {RES['fail']} fallidas")
    if FALLOS:
        print("\nFALLOS:")
        for f_ in FALLOS:
            print(" ", f_)
    print("=" * 90)
    return 0 if RES["fail"] == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
