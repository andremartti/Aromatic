# -*- coding: utf-8 -*-
"""Estilos, formatos y utilidades comunes del sistema financiero AROMATIC."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- Paleta ----
AZUL_OSCURO = "1F3A5F"   # encabezados principales
AZUL_MEDIO = "2E5C8A"    # encabezados secundarios
GRIS_BORDE = "C7CDD4"
FONDO_INPUT = "FFF8E1"   # celdas de captura (amarillo tenue)
FONDO_CALC = "F4F6F9"    # celdas calculadas
FONDO_TITULO = "EAF0F6"
BLANCO = "FFFFFF"
VERDE = "1E7B4F"
ROJO = "B00020"
AMBAR = "8A6100"

# --------------------------------------------------------------- Formatos ---
FMT_MONEDA = '"L"#,##0.00;[Red]-"L"#,##0.00;"L"0.00'
FMT_MONEDA0 = '"L"#,##0;[Red]-"L"#,##0;"L"0'
FMT_PCT = '0.0%;[Red]-0.0%;0.0%'
FMT_PCT2 = '0.00%'
FMT_NUM = '#,##0;[Red]-#,##0;0'
FMT_NUM2 = '#,##0.00'
FMT_FECHA = 'dd/mm/yyyy'
FMT_MES = 'mmm-yyyy'
FMT_TEXTO = '@'

# ----------------------------------------------------------------- Fuentes --
F_TITULO = Font(name="Calibri", size=16, bold=True, color=AZUL_OSCURO)
F_SUBTITULO = Font(name="Calibri", size=12, bold=True, color=AZUL_MEDIO)
F_HEADER = Font(name="Calibri", size=10, bold=True, color=BLANCO)
F_NORMAL = Font(name="Calibri", size=10)
F_BOLD = Font(name="Calibri", size=10, bold=True)
F_NOTA = Font(name="Calibri", size=9, italic=True, color="5A6671")
F_KPI = Font(name="Calibri", size=18, bold=True, color=AZUL_OSCURO)
F_KPI_LABEL = Font(name="Calibri", size=9, bold=True, color="5A6671")

# ------------------------------------------------------------------ Rellenos -
FILL_HEADER = PatternFill("solid", fgColor=AZUL_OSCURO)
FILL_HEADER2 = PatternFill("solid", fgColor=AZUL_MEDIO)
FILL_INPUT = PatternFill("solid", fgColor=FONDO_INPUT)
FILL_CALC = PatternFill("solid", fgColor=FONDO_CALC)
FILL_TITULO = PatternFill("solid", fgColor=FONDO_TITULO)
FILL_BLANCO = PatternFill("solid", fgColor=BLANCO)

# ------------------------------------------------------------------ Bordes ---
_lado = Side(style="thin", color=GRIS_BORDE)
BORDE = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)
_lado_grueso = Side(style="medium", color=AZUL_OSCURO)
BORDE_KPI = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)

AL_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_IZQ = Alignment(horizontal="left", vertical="center")
AL_IZQ_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_DER = Alignment(horizontal="right", vertical="center")

DESBLOQUEADA = Protection(locked=False)
BLOQUEADA = Protection(locked=True)


def titulo(ws, celda, texto, subtitulo=None):
    """Escribe el título de una hoja."""
    ws[celda] = texto
    ws[celda].font = F_TITULO
    if subtitulo:
        fila = int("".join(c for c in celda if c.isdigit())) + 1
        col = "".join(c for c in celda if c.isalpha())
        ws[f"{col}{fila}"] = subtitulo
        ws[f"{col}{fila}"].font = F_NOTA


def encabezado_tabla(ws, fila, col_ini, encabezados, alto=32):
    """Pinta la fila de encabezados de una tabla."""
    for i, h in enumerate(encabezados):
        c = ws.cell(row=fila, column=col_ini + i, value=h)
        c.font = F_HEADER
        c.fill = FILL_HEADER
        c.alignment = AL_CENTRO
        c.border = BORDE
    ws.row_dimensions[fila].height = alto


def banda(ws, fila, col_ini, col_fin, texto, fill=None):
    """Banda/sección de color con un rótulo."""
    ws.cell(row=fila, column=col_ini, value=texto).font = F_SUBTITULO
    for c in range(col_ini, col_fin + 1):
        ws.cell(row=fila, column=c).fill = fill or FILL_TITULO


def anchos(ws, mapa):
    for col, w in mapa.items():
        ws.column_dimensions[col].width = w


def marcar_input(celda):
    celda.fill = FILL_INPUT
    celda.protection = DESBLOQUEADA
    celda.border = BORDE
    if celda.font is None or not celda.font.bold:
        celda.font = F_NORMAL


def marcar_calc(celda):
    celda.fill = FILL_CALC
    celda.protection = BLOQUEADA
    celda.border = BORDE
    celda.font = F_NORMAL


def col(n):
    return get_column_letter(n)
